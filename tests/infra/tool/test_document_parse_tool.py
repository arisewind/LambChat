"""document_parse 工具与多提供商层测试。

提供商契约对齐 Open WebUI 的文档解析接入面：
Mistral OCR / MinerU（cloud+local）/ Azure Document Intelligence /
docling-serve / Apache Tika / PaddleOCR-VL。
"""

from __future__ import annotations

import base64
import io
import json
import zipfile
from types import SimpleNamespace

import pytest


class _Runtime:
    def __init__(self, user_id: str | None, base_url: str = "https://app.example.com") -> None:
        context = SimpleNamespace(user_id=user_id) if user_id is not None else None
        self.config = {"configurable": {"context": context, "base_url": base_url}}


def _fake_ocr_payload() -> dict:
    return {
        "model": "mistral-ocr-latest",
        "pages": [
            {
                "index": 1,
                "markdown": "# Hello\n\n![image](img-0.jpeg)",
                "images": [
                    {"id": "img-0.jpeg", "image_base64": base64.b64encode(b"PNGDATA").decode()}
                ],
            },
            {"index": 2, "markdown": "Second page"},
        ],
        "usage_info": {"pages_processed": 2},
    }


# ── 纯函数 ────────────────────────────────────────────────────────────────


def test_build_mistral_ocr_request_defaults() -> None:
    from src.infra.tool.document_parse_providers import build_mistral_ocr_request

    body = build_mistral_ocr_request("mistral-ocr-latest", "report.docx", "QkFTRTY0")

    assert body["model"] == "mistral-ocr-latest"
    assert body["document"] == {
        "type": "base64",
        "base64": "QkFTRTY0",
        "document_name": "report.docx",
    }
    assert body["include_image_base64"] is True


def test_build_mistral_ocr_request_with_pages_and_image_limit() -> None:
    from src.infra.tool.document_parse_providers import build_mistral_ocr_request

    body = build_mistral_ocr_request(
        "ocr-4",
        "report.pdf",
        "QkFTRTY0",
        include_images=False,
        image_limit=5,
        pages="0-2",
    )

    assert "include_image_base64" not in body
    assert body["pages"] == "0-2"
    # include_images=False 时不带图片限制
    assert "image_limit" not in body


def test_merge_page_markdown_sorts_by_index() -> None:
    from src.infra.tool.document_parse_providers import merge_page_markdown

    pages = [
        {"index": 2, "markdown": "Second"},
        {"index": 1, "markdown": "First"},
    ]

    assert merge_page_markdown(pages) == "First\n\nSecond"


def test_extract_ocr_images_dedupes_and_orders() -> None:
    from src.infra.tool.document_parse_providers import extract_ocr_images

    pages = [
        {
            "index": 1,
            "markdown": "x",
            "images": [
                {"id": "img-0.jpeg", "image_base64": "QQ=="},
                {"id": "img-1.png", "image_base64": "Qg=="},
            ],
        },
        {"index": 2, "markdown": "y", "images": [{"id": "img-0.jpeg", "image_base64": "QQ=="}]},
    ]

    images = extract_ocr_images(pages)

    assert [image["ref"] for image in images] == ["img-0.jpeg", "img-1.png"]


def test_extract_mineru_zip_markdown_and_images() -> None:
    from src.infra.tool.document_parse_providers import extract_mineru_zip

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("doc/full/abc123.md", "# MinerU\n\n![f](images/fig1.jpg)")
        archive.writestr("doc/full/images/fig1.jpg", "JPGDATA")
        archive.writestr("doc/full/layout.pdf", "not-an-image")
    zip_bytes = buffer.getvalue()

    markdown, images = extract_mineru_zip(zip_bytes)

    assert markdown == "# MinerU\n\n![f](images/fig1.jpg)"
    assert images == [{"ref": "images/fig1.jpg", "base64": base64.b64encode(b"JPGDATA").decode()}]


def test_rewrite_image_refs_replaces_only_known_refs() -> None:
    from src.infra.tool.document_parse_tool import rewrite_image_refs

    markdown = "![a](img-0.jpeg) ![b](images/fig1.jpg) [link](https://example.com)"

    rewritten = rewrite_image_refs(
        markdown,
        {"img-0.jpeg": "https://files/1.jpeg", "images/fig1.jpg": "https://files/fig1.jpg"},
    )

    assert "![a](https://files/1.jpeg)" in rewritten
    assert "![b](https://files/fig1.jpg)" in rewritten
    assert "[link](https://example.com)" in rewritten


def test_truncate_markdown_marks_truncated() -> None:
    from src.infra.tool.document_parse_tool import truncate_markdown

    kept, truncated = truncate_markdown("x" * 100, 40)

    assert truncated is True
    assert kept.startswith("x")
    assert len(kept) <= 40
    assert "[truncated]" in kept

    same, not_truncated = truncate_markdown("abc", 10)
    assert same == "abc"
    assert not_truncated is False


def test_detect_document_extension() -> None:
    from src.infra.tool.document_parse_tool import detect_document_extension

    assert detect_document_extension("https://x/api/upload/file/a/report.PDF") == "pdf"
    assert detect_document_extension("doc.docx") == "docx"
    assert detect_document_extension("slides.pptx") == "pptx"
    assert detect_document_extension("legacy.doc") == "doc"
    assert detect_document_extension("book.odt") == "odt"
    assert detect_document_extension("sheet.xlsx") is None
    assert detect_document_extension("photo.png") is None


def test_mime_type_from_image_id() -> None:
    from src.infra.tool.document_parse_tool import mime_type_from_image_id

    assert mime_type_from_image_id("img-0.jpeg") == "image/jpeg"
    assert mime_type_from_image_id("img-1.png") == "image/png"
    assert mime_type_from_image_id("img-2.bin") == "application/octet-stream"


# ── provider 链选择 ───────────────────────────────────────────────────────

_PROVIDER_SETTING_NAMES = (
    "DOCUMENT_PARSE_PROVIDER",
    "DOCUMENT_PARSE_MISTRAL_API_KEY",
    "DOCUMENT_PARSE_MISTRAL_BASE_URL",
    "DOCUMENT_PARSE_MISTRAL_MODEL",
    "DOCUMENT_PARSE_MINERU_API_MODE",
    "DOCUMENT_PARSE_MINERU_API_URL",
    "DOCUMENT_PARSE_MINERU_API_KEY",
    "DOCUMENT_PARSE_AZURE_ENDPOINT",
    "DOCUMENT_PARSE_AZURE_KEY",
    "DOCUMENT_PARSE_AZURE_MODEL",
    "DOCUMENT_PARSE_DOCLING_URL",
    "DOCUMENT_PARSE_DOCLING_API_KEY",
    "DOCUMENT_PARSE_PADDLEOCR_URL",
    "DOCUMENT_PARSE_PADDLEOCR_TOKEN",
    "DOCUMENT_PARSE_TIKA_URL",
)


def _clear_provider_settings(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    for name in _PROVIDER_SETTING_NAMES:
        monkeypatch.setattr(providers.settings, name, "")


def test_provider_chain_auto_picks_first_configured(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    # markitdown 是本地库，零配置即可用 → 链尾兜底始终存在
    assert providers.resolve_document_parse_provider_chain() == ["markitdown"]

    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_DOCLING_URL", "http://docling:5001")
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MISTRAL_API_KEY", "sk-mistral")
    # auto 顺序：mistral 优先于 docling，markitdown 垫底
    assert providers.resolve_document_parse_provider_chain() == [
        "mistral",
        "docling",
        "markitdown",
    ]


def test_provider_chain_empty_when_markitdown_missing(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers, "_markitdown_available", False)
    assert providers.resolve_document_parse_provider_chain() == []


def test_provider_chain_pinned_provider(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_PROVIDER", "tika")
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_TIKA_URL", "http://tika:9998")
    assert providers.resolve_document_parse_provider_chain() == ["tika"]


def test_mineru_cloud_requires_key_local_requires_url(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MINERU_API_KEY", "mk")
    assert providers._provider_available("mineru") is True

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MINERU_API_MODE", "local")
    monkeypatch.setattr(
        providers.settings, "DOCUMENT_PARSE_MINERU_API_URL", "http://localhost:8000"
    )
    assert providers._provider_available("mineru") is True


# ── provider 请求构造与响应解析 ───────────────────────────────────────────


class _FakeResponse:
    def __init__(
        self,
        payload: dict | list | Exception | None = None,
        *,
        headers: dict[str, str] | None = None,
        content: bytes = b"",
    ) -> None:
        self._payload = payload
        self.headers = headers or {}
        self.content = content

    def raise_for_status(self) -> None:
        if isinstance(self._payload, Exception):
            raise self._payload

    @property
    def is_success(self) -> bool:
        return not isinstance(self._payload, Exception)

    def json(self) -> dict | list:
        if isinstance(self._payload, Exception):
            raise self._payload
        return self._payload  # type: ignore[return-value]


class _FakeClient:
    """按请求顺序回放响应的假 httpx.AsyncClient。"""

    def __init__(self, responses: list[_FakeResponse]) -> None:
        self._responses = list(responses)
        self.calls: list[dict] = []

    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, tb):
        return None

    def _next(self, method: str, url: str, **kwargs) -> _FakeResponse:
        self.calls.append({"method": method, "url": url, **kwargs})
        if not self._responses:
            raise AssertionError(f"unexpected request: {method} {url}")
        return self._responses.pop(0)

    async def post(self, url: str, **kwargs) -> _FakeResponse:
        return self._next("POST", url, **kwargs)

    async def put(self, url: str, **kwargs) -> _FakeResponse:
        return self._next("PUT", url, **kwargs)

    async def get(self, url: str, **kwargs) -> _FakeResponse:
        return self._next("GET", url, **kwargs)


async def _no_sleep(_: float) -> None:
    return None


@pytest.mark.asyncio
async def test_mistral_parse_builds_request_and_extracts_images(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MISTRAL_API_KEY", "sk-mistral")
    monkeypatch.setattr(
        providers.settings, "DOCUMENT_PARSE_MISTRAL_BASE_URL", "https://api.example.com"
    )

    client = _FakeClient([_FakeResponse(_fake_ocr_payload())])
    result = await providers.mistral_parse(
        client,
        data=b"PDFBYTES",
        filename="report.docx",
        include_images=True,
        pages="0-1",
        image_limit=5,
    )

    assert result["markdown"] == "# Hello\n\n![image](img-0.jpeg)\n\nSecond page"
    assert result["images"] == [
        {"ref": "img-0.jpeg", "base64": base64.b64encode(b"PNGDATA").decode()}
    ]
    assert result["pages"] == 2
    assert result["engine"] == "mistral:mistral-ocr-latest"

    call = client.calls[0]
    assert call["url"] == "https://api.example.com/v1/ocr"
    assert call["headers"] == {"Authorization": "Bearer sk-mistral"}
    assert call["json"]["document"]["base64"] == base64.b64encode(b"PDFBYTES").decode()
    assert call["json"]["document"]["document_name"] == "report.docx"
    assert call["json"]["pages"] == "0-1"
    assert call["json"]["image_limit"] == 5


@pytest.mark.asyncio
async def test_mineru_cloud_parse_full_flow(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MINERU_API_KEY", "mk")
    monkeypatch.setattr(providers.asyncio, "sleep", _no_sleep)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("doc/full/abc.md", "# MinerU\n\n![f](images/fig1.jpg)")
        archive.writestr("doc/full/images/fig1.jpg", "JPGDATA")

    client = _FakeClient(
        [
            _FakeResponse(
                {
                    "code": 0,
                    "data": {
                        "batch_id": "batch-1",
                        "file_urls": ["https://cdn/upload-presigned"],
                    },
                }
            ),
            _FakeResponse({}),
            _FakeResponse(
                {
                    "code": 0,
                    "data": {
                        "extract_result": [
                            {
                                "file_name": "report.pdf",
                                "state": "done",
                                "full_zip_url": "https://cdn/result.zip",
                            }
                        ]
                    },
                }
            ),
            _FakeResponse(content=buffer.getvalue()),
        ]
    )
    result = await providers.mineru_parse(
        client,
        data=b"PDFBYTES",
        filename="report.pdf",
        include_images=True,
        pages=None,
        image_limit=0,
    )

    assert result["engine"] == "mineru:cloud"
    assert result["markdown"] == "# MinerU\n\n![f](images/fig1.jpg)"
    assert result["images"] == [
        {"ref": "images/fig1.jpg", "base64": base64.b64encode(b"JPGDATA").decode()}
    ]
    assert client.calls[0]["url"] == "https://mineru.net/api/v4/file-urls/batch"
    assert client.calls[1]["url"] == "https://cdn/upload-presigned"
    assert client.calls[2]["url"] == "https://mineru.net/api/v4/extract-results/batch/batch-1"
    assert client.calls[3]["url"] == "https://cdn/result.zip"


@pytest.mark.asyncio
async def test_mineru_local_parse(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MINERU_API_MODE", "local")
    monkeypatch.setattr(
        providers.settings, "DOCUMENT_PARSE_MINERU_API_URL", "http://localhost:8000"
    )

    payload = {
        "results": {
            "report.pdf": {
                "md_content": "# Local\n\n![i](img-1.png)",
                "images": {"img-1.png": base64.b64encode(b"LOCALPNG").decode()},
            }
        }
    }
    client = _FakeClient([_FakeResponse(payload)])
    result = await providers.mineru_parse(
        client,
        data=b"PDFBYTES",
        filename="report.pdf",
        include_images=True,
        pages=None,
        image_limit=0,
    )

    assert result["engine"] == "mineru:local"
    assert result["markdown"] == "# Local\n\n![i](img-1.png)"
    assert result["images"] == [
        {"ref": "img-1.png", "base64": base64.b64encode(b"LOCALPNG").decode()}
    ]
    assert client.calls[0]["url"] == "http://localhost:8000/file_parse"


@pytest.mark.asyncio
async def test_azure_parse_polls_operation(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(
        providers.settings, "DOCUMENT_PARSE_AZURE_ENDPOINT", "https://di.example.com"
    )
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_AZURE_KEY", "az-key")
    monkeypatch.setattr(providers.asyncio, "sleep", _no_sleep)

    client = _FakeClient(
        [
            _FakeResponse(
                {},
                headers={"operation-location": "https://di.example.com/operations/op-1"},
            ),
            _FakeResponse({"status": "running"}),
            _FakeResponse(
                {
                    "status": "succeeded",
                    "analyzeResult": {"content": "# Azure\n\nBody", "pages": [{}, {}, {}]},
                }
            ),
        ]
    )
    result = await providers.azure_parse(
        client,
        data=b"PDFBYTES",
        filename="report.docx",
        include_images=True,
        pages=None,
        image_limit=0,
    )

    assert result["markdown"] == "# Azure\n\nBody"
    assert result["pages"] == 3
    assert result["engine"] == "azure:prebuilt-read"
    assert client.calls[0]["url"].endswith(
        "/documentintelligence/documentModels/prebuilt-read:analyze"
    )
    assert client.calls[0]["params"]["outputContentFormat"] == "markdown"
    assert client.calls[0]["headers"] == {"Ocp-Apim-Subscription-Key": "az-key"}
    assert client.calls[1]["url"] == "https://di.example.com/operations/op-1"


@pytest.mark.asyncio
async def test_docling_parse(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_DOCLING_URL", "http://docling:5001")
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_DOCLING_API_KEY", "dk")

    client = _FakeClient([_FakeResponse({"document": {"md": "# Docling\n\nOK", "num_pages": 4}})])
    result = await providers.docling_parse(
        client,
        data=b"PDFBYTES",
        filename="report.pdf",
        include_images=True,
        pages=None,
        image_limit=0,
    )

    assert result["markdown"] == "# Docling\n\nOK"
    assert result["pages"] == 4
    assert result["engine"] == "docling"
    assert client.calls[0]["url"] == "http://docling:5001/v1/convert/file"
    assert client.calls[0]["headers"] == {"X-Api-Key": "dk"}
    assert client.calls[0]["files"]["files"] == (
        "report.pdf",
        b"PDFBYTES",
        "application/octet-stream",
    )


@pytest.mark.asyncio
async def test_tika_parse(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_TIKA_URL", "http://tika:9998")

    client = _FakeClient([_FakeResponse({"X-TIKA:content": "Plain text from tika"})])
    result = await providers.tika_parse(
        client,
        data=b"OLEDOC",
        filename="legacy.doc",
        include_images=True,
        pages=None,
        image_limit=0,
    )

    assert result["markdown"] == "Plain text from tika"
    assert result["images"] == []
    assert result["engine"] == "tika"
    assert client.calls[0]["method"] == "PUT"
    assert client.calls[0]["url"] == "http://tika:9998/tika/text"


@pytest.mark.asyncio
async def test_paddleocr_parse(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_PADDLEOCR_URL", "http://paddle:8080")
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_PADDLEOCR_TOKEN", "pt")

    payload = {
        "result": {
            "layoutParsingResults": [
                {
                    "markdown": {
                        "text": "# Paddle\n\n![p](p0.png)",
                        "images": {"p0.png": base64.b64encode(b"PADDLEPNG").decode()},
                    }
                }
            ]
        }
    }
    client = _FakeClient([_FakeResponse(payload)])
    result = await providers.paddleocr_parse(
        client,
        data=b"PDFBYTES",
        filename="report.pdf",
        include_images=True,
        pages=None,
        image_limit=0,
    )

    assert result["markdown"] == "# Paddle\n\n![p](p0.png)"
    assert result["images"] == [
        {"ref": "p0.png", "base64": base64.b64encode(b"PADDLEPNG").decode()}
    ]
    assert result["pages"] == 1
    assert result["engine"] == "paddleocr_vl"
    assert client.calls[0]["headers"]["Authorization"] == "token pt"
    assert client.calls[0]["json"]["file"] == base64.b64encode(b"PDFBYTES").decode()
    assert client.calls[0]["json"]["fileType"] == 0


# ── MarkItDown（本地库） ──────────────────────────────────────────────────

_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _build_docx_with_image() -> bytes:
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
        ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        "<w:body><w:p><w:r><w:t>Quarterly Report</w:t></w:r></w:p>"
        '<w:p><w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="9525" cy="9525"/><wp:docPr id="1" name="Picture 1"/>'
        "<a:graphic><a:graphicData"
        ' uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic><pic:nvPicPr><pic:cNvPr id="1"'
        ' name="image1.png"/><pic:cNvPicPr/></pic:nvPicPr>'
        '<pic:blipFill><a:blip r:embed="rId7"/></pic:blipFill>'
        '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="9525"'
        ' cy="9525"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
        "</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p></w:body></w:document>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n'
        '<Relationship Id="rId7"'
        ' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"'
        ' Target="media/image1.png"/>\n'
        "</Relationships>"
    )
    ctypes = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels"'
        ' ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Default Extension="png" ContentType="image/png"/>'
        '<Override PartName="/word/document.xml"'
        ' ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", ctypes)
        archive.writestr("word/document.xml", document)
        archive.writestr("word/_rels/document.xml.rels", rels)
        archive.writestr("word/media/image1.png", _TINY_PNG)
    return buffer.getvalue()


def _build_pptx_with_image() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 100, 20)
    box.text_frame.text = "Slide one text"
    slide.shapes.add_picture(io.BytesIO(_TINY_PNG), 0, 30)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def test_extract_ooxml_media_images() -> None:
    from src.infra.tool.document_parse_providers import extract_ooxml_media_images

    docx = _build_docx_with_image()
    images = extract_ooxml_media_images(docx)

    assert [image["ref"] for image in images] == ["markitdown-img-1.png"]
    assert images[0]["base64"] == base64.b64encode(_TINY_PNG).decode()

    assert extract_ooxml_media_images(b"not a zip") == []


def test_rewrite_local_image_refs_sequential_and_skips_http() -> None:
    from src.infra.tool.document_parse_providers import rewrite_local_image_refs

    markdown = "![a](data:image/png;base64...) [link](https://example.com) ![b](Picture2.jpg)"

    rewritten = rewrite_local_image_refs(markdown, ["markitdown-img-1.png", "markitdown-img-2.jpg"])

    assert "![a](markitdown-img-1.png)" in rewritten
    assert "![b](markitdown-img-2.jpg)" in rewritten
    assert "[link](https://example.com)" in rewritten
    # refs 用尽后剩余引用原样保留
    tail = rewrite_local_image_refs("![x](a.png) ![y](b.png)", ["only-1.png"])
    assert "![x](only-1.png)" in tail
    assert "![y](b.png)" in tail


@pytest.mark.asyncio
async def test_markitdown_parse_docx_with_image() -> None:
    from src.infra.tool import document_parse_providers as providers

    class _UnusedClient:
        pass

    result = await providers.markitdown_parse(
        _UnusedClient(),  # type: ignore[arg-type]
        data=_build_docx_with_image(),
        filename="report.docx",
        include_images=True,
        pages=None,
        image_limit=0,
    )

    assert result["engine"] == "markitdown"
    assert "Quarterly Report" in result["markdown"]
    # 截断的 data URI 占位符被换成生成的 ref，图片字节来自 ZIP media
    assert "data:image" not in result["markdown"]
    assert result["images"] == [
        {"ref": "markitdown-img-1.png", "base64": base64.b64encode(_TINY_PNG).decode()}
    ]
    assert "](markitdown-img-1.png)" in result["markdown"]


@pytest.mark.asyncio
async def test_markitdown_parse_pptx_with_image() -> None:
    from src.infra.tool import document_parse_providers as providers

    class _UnusedClient:
        pass

    result = await providers.markitdown_parse(
        _UnusedClient(),  # type: ignore[arg-type]
        data=_build_pptx_with_image(),
        filename="slides.pptx",
        include_images=True,
        pages=None,
        image_limit=0,
    )

    assert result["engine"] == "markitdown"
    assert "Slide one text" in result["markdown"]
    assert [image["ref"] for image in result["images"]] == ["markitdown-img-1.png"]
    assert "](markitdown-img-1.png)" in result["markdown"]


@pytest.mark.asyncio
async def test_markitdown_parse_rejects_unsupported_extension() -> None:
    from src.infra.tool import document_parse_providers as providers

    class _UnusedClient:
        pass

    with pytest.raises(providers.DocumentParseError) as exc_info:
        await providers.markitdown_parse(
            _UnusedClient(),  # type: ignore[arg-type]
            data=b"OLE",
            filename="legacy.doc",
            include_images=True,
            pages=None,
            image_limit=0,
        )

    assert "markitdown does not support" in str(exc_info.value)


# ── 纯图片文档自动转投 OCR（docx/pptx 整页贴图，文本层为空） ────────────


def test_markdown_text_content_strips_image_refs_only() -> None:
    from src.infra.tool.document_parse_providers import markdown_text_content

    image_only = "![](markitdown-img-1.png)\n\n![](markitdown-img-2.png)\n"
    assert markdown_text_content(image_only) == ""
    mixed = "# 标题\n\n![](http://x/a.png)\n\n正文段落"
    stripped = markdown_text_content(mixed)
    # 只剥图片引用本身，保留其余结构（含周边空行）
    assert "![" not in stripped
    assert stripped.replace("\n", "") == "# 标题正文段落"
    table = "<table><tr><td>signal</td></tr></table>"
    assert markdown_text_content(table) == table


@pytest.mark.asyncio
async def test_execute_document_parse_redispatches_image_only_result(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MISTRAL_API_KEY", "sk-mistral")

    calls: list[str] = []

    async def _branching_parse(client, *, data, filename, include_images, pages, image_limit):
        calls.append(filename)
        if filename.endswith((".docx", ".pptx")):
            # 首遍：文本层为空，只有页面图
            return providers._result(
                markdown="![](markitdown-img-1.png)\n\n![](markitdown-img-2.png)",
                images=[
                    {"ref": "markitdown-img-1.png", "base64": base64.b64encode(b"P1").decode()},
                    {"ref": "markitdown-img-2.png", "base64": base64.b64encode(b"P2").decode()},
                ],
                pages=None,
                engine="mistral:mistral-ocr-latest",
            )
        index = filename.removeprefix("page-").removesuffix(".png")
        return providers._result(
            markdown=f"## Page {index} OCR 文本", images=[], pages=1, engine="mistral:ocr"
        )

    monkeypatch.setitem(providers.PROVIDER_FUNCS, "mistral", _branching_parse)

    result = await providers.execute_document_parse(
        data=b"DOCX",
        filename="spec.docx",
        include_images=True,
        client=_FakeClient([]),
    )

    assert calls == ["spec.docx", "page-1.png", "page-2.png"]
    assert result["engine"] == "mistral:mistral-ocr-latest+ocr:mistral"
    # 每页 OCR 文本紧跟其原图引用（保住 VLM 视觉兜底）
    assert "## Page 1 OCR 文本" in result["markdown"]
    assert "## Page 2 OCR 文本" in result["markdown"]
    assert "![](markitdown-img-1.png)" in result["markdown"]
    assert "![](markitdown-img-2.png)" in result["markdown"]
    assert result["pages"] == 2


@pytest.mark.asyncio
async def test_execute_document_parse_skips_redispatch_when_text_present(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MISTRAL_API_KEY", "sk-mistral")

    calls: list[str] = []

    async def _text_parse(client, *, data, filename, include_images, pages, image_limit):
        calls.append(filename)
        return providers._result(
            markdown="# 有文字的文档\n\n![](markitdown-img-1.png)",
            images=[{"ref": "markitdown-img-1.png", "base64": "UE5H"}],
            pages=None,
            engine="mistral:mistral-ocr-latest",
        )

    monkeypatch.setitem(providers.PROVIDER_FUNCS, "mistral", _text_parse)

    result = await providers.execute_document_parse(
        data=b"DOCX",
        filename="spec.docx",
        include_images=True,
        client=_FakeClient([]),
    )

    assert calls == ["spec.docx"]
    assert result["engine"] == "mistral:mistral-ocr-latest"
    assert result["markdown"] == "# 有文字的文档\n\n![](markitdown-img-1.png)"


@pytest.mark.asyncio
async def test_execute_document_parse_skips_redispatch_without_capable_provider(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers, "_markitdown_available", True)

    calls: list[str] = []

    async def _image_only_parse(client, *, data, filename, include_images, pages, image_limit):
        calls.append(filename)
        return providers._result(
            markdown="![](markitdown-img-1.png)",
            images=[{"ref": "markitdown-img-1.png", "base64": "UE5H"}],
            pages=None,
            engine="markitdown",
        )

    monkeypatch.setitem(providers.PROVIDER_FUNCS, "markitdown", _image_only_parse)

    result = await providers.execute_document_parse(
        data=b"DOCX",
        filename="spec.docx",
        include_images=True,
        client=_FakeClient([]),
    )

    # 链上只有 markitdown（无 OCR 能力）：原样返回，不重投
    assert calls == ["spec.docx"]
    assert result["engine"] == "markitdown"
    assert result["markdown"] == "![](markitdown-img-1.png)"


@pytest.mark.asyncio
async def test_execute_document_parse_partial_redispatch_failure_keeps_image_refs(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MISTRAL_API_KEY", "sk-mistral")

    async def _branching_parse(client, *, data, filename, include_images, pages, image_limit):
        if filename.endswith((".docx", ".pptx")):
            return providers._result(
                markdown="![](markitdown-img-1.png)\n\n![](markitdown-img-2.png)",
                images=[
                    {"ref": "markitdown-img-1.png", "base64": "UTE="},
                    {"ref": "markitdown-img-2.png", "base64": "UTI="},
                ],
                pages=None,
                engine="mistral:mistral-ocr-latest",
            )
        if filename == "page-2.png":
            raise providers.DocumentParseError("mistral HTTP 500")
        return providers._result(
            markdown="第一页 OCR 文本", images=[], pages=1, engine="mistral:ocr"
        )

    monkeypatch.setitem(providers.PROVIDER_FUNCS, "mistral", _branching_parse)

    result = await providers.execute_document_parse(
        data=b"DOCX",
        filename="spec.docx",
        include_images=True,
        client=_FakeClient([]),
    )

    # page-1 出字、page-2 失败保图：合并结果两页图引用都在；pages 记文档总页数
    assert "第一页 OCR 文本" in result["markdown"]
    assert "![](markitdown-img-1.png)" in result["markdown"]
    assert "![](markitdown-img-2.png)" in result["markdown"]
    assert result["pages"] == 2


@pytest.mark.asyncio
async def test_execute_document_parse_all_redispatch_failure_returns_original(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MISTRAL_API_KEY", "sk-mistral")

    async def _branching_parse(client, *, data, filename, include_images, pages, image_limit):
        if not filename.endswith(".png"):
            return providers._result(
                markdown="![](markitdown-img-1.png)",
                images=[{"ref": "markitdown-img-1.png", "base64": "UTE="}],
                pages=None,
                engine="mistral:mistral-ocr-latest",
            )
        raise providers.DocumentParseError("mistral HTTP 503")

    monkeypatch.setitem(providers.PROVIDER_FUNCS, "mistral", _branching_parse)

    result = await providers.execute_document_parse(
        data=b"DOCX",
        filename="spec.docx",
        include_images=True,
        client=_FakeClient([]),
    )

    # 全部页面 OCR 失败：保留首遍原结果（引擎与 markdown 不变）
    assert result["engine"] == "mistral:mistral-ocr-latest"
    assert result["markdown"] == "![](markitdown-img-1.png)"


@pytest.mark.asyncio
async def test_execute_document_parse_no_provider(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers, "_markitdown_available", False)

    with pytest.raises(providers.DocumentParseError) as exc_info:
        await providers.execute_document_parse(data=b"x", filename="a.pdf", include_images=True)

    assert str(exc_info.value) == "document_parse_no_provider_configured"


@pytest.mark.asyncio
async def test_execute_document_parse_falls_back_to_next_provider(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_providers as providers

    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_MISTRAL_API_KEY", "sk-mistral")
    monkeypatch.setattr(providers.settings, "DOCUMENT_PARSE_TIKA_URL", "http://tika:9998")

    async def _failing_parse(*args, **kwargs):
        raise providers.DocumentParseError("mistral HTTP 429: quota")

    monkeypatch.setitem(providers.PROVIDER_FUNCS, "mistral", _failing_parse)

    tika_client = _FakeClient([_FakeResponse({"X-TIKA:content": "fallback text"})])
    result = await providers.execute_document_parse(
        data=b"PDFBYTES",
        filename="a.pdf",
        include_images=True,
        client=tika_client,
    )

    assert result["engine"] == "tika"
    assert result["markdown"] == "fallback text"


# ── 工具级异步链路 ────────────────────────────────────────────────────────


def _patch_common(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.infra.tool import document_parse_tool

    async def fake_run_long_blocking_io(func, *args, **kwargs):
        return func(*args, **kwargs)

    monkeypatch.setattr(document_parse_tool, "run_long_blocking_io", fake_run_long_blocking_io)
    _clear_provider_settings(monkeypatch)
    monkeypatch.setattr(document_parse_tool.settings, "DOCUMENT_PARSE_MISTRAL_API_KEY", "sk-ocr")
    monkeypatch.setattr(
        document_parse_tool.settings, "DOCUMENT_PARSE_MISTRAL_BASE_URL", "https://api.example.com"
    )
    monkeypatch.setattr(document_parse_tool.settings, "DOCUMENT_PARSE_MAX_DOWNLOAD_BYTES", 52428800)
    monkeypatch.setattr(document_parse_tool.settings, "DOCUMENT_PARSE_MAX_OUTPUT_CHARS", 200000)
    monkeypatch.setattr(document_parse_tool.settings, "DOCUMENT_PARSE_MAX_IMAGES", 20)


class _FakeStreamResponse:
    def __init__(self, chunks: list[bytes], headers: dict[str, str] | None = None) -> None:
        self._chunks = chunks
        self.headers = headers or {}

    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, tb):
        return None

    def raise_for_status(self) -> None:
        return None

    async def aiter_bytes(self):
        for chunk in self._chunks:
            yield chunk


class _FakeToolHttpClient:
    """document_parse 工具用的假客户端：stream 下载 + post 转发 OCR。"""

    def __init__(self, chunks: list[bytes], ocr_payload: dict | Exception) -> None:
        self._chunks = chunks
        self._ocr_payload = ocr_payload
        self.download_urls: list[str] = []
        self.post_calls: list[dict] = []

    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, tb):
        return None

    def stream(self, method: str, request_url: str):
        assert method == "GET"
        self.download_urls.append(request_url)
        return _FakeStreamResponse(self._chunks)

    async def post(self, request_url: str, **kwargs):
        self.post_calls.append({"url": request_url, **kwargs})
        if isinstance(self._ocr_payload, Exception):
            raise self._ocr_payload
        return _FakeResponse(self._ocr_payload)


def _patch_http(monkeypatch: pytest.MonkeyPatch, client: _FakeToolHttpClient) -> None:
    from src.infra.tool import document_parse_tool

    monkeypatch.setattr(document_parse_tool.httpx, "AsyncClient", lambda **kwargs: client)


def _patch_image_upload(monkeypatch: pytest.MonkeyPatch) -> list[dict]:
    from src.infra.tool import document_parse_tool

    uploaded: list[dict] = []

    async def fake_upload(*, user_id: str, image_id: str, data: bytes):
        uploaded.append({"user_id": user_id, "image_id": image_id, "data": data})
        return {
            "url": f"https://storage.example.com/parsed-documents/{user_id}/{image_id}",
            "key": f"parsed-documents/{user_id}/{image_id}",
            "size": len(data),
            "content_type": "image/jpeg",
        }

    monkeypatch.setattr(document_parse_tool, "_upload_image_to_storage", fake_upload)
    return uploaded


@pytest.mark.asyncio
async def test_document_parse_parses_document_and_rewrites_images(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_tool

    _patch_common(monkeypatch)
    client = _FakeToolHttpClient([b"PDF", b"BYTES"], _fake_ocr_payload())
    _patch_http(monkeypatch, client)
    uploaded = _patch_image_upload(monkeypatch)

    result = json.loads(
        await document_parse_tool.document_parse.coroutine(
            url="/api/upload/file/docs/report.docx",
            include_images=True,
            runtime=_Runtime("user-1"),
        )
    )

    assert result["success"] is True
    assert result["filename"] == "report.docx"
    assert result["url"] == "https://app.example.com/api/upload/file/docs/report.docx"
    assert result["provider"] == "mistral:mistral-ocr-latest"
    assert result["pages_processed"] == 2
    assert result["truncated"] is False
    assert client.download_urls == ["https://app.example.com/api/upload/file/docs/report.docx"]
    assert client.post_calls[0]["url"] == "https://api.example.com/v1/ocr"
    request_body = client.post_calls[0]["json"]
    assert request_body["document"]["base64"] == base64.b64encode(b"PDFBYTES").decode()
    assert request_body["document"]["document_name"] == "report.docx"
    assert request_body["include_image_base64"] is True
    assert result["markdown"] == (
        "# Hello\n\n"
        "![image](https://storage.example.com/parsed-documents/user-1/img-0.jpeg)\n\n"
        "Second page"
    )
    assert result["image_count"] == 1
    assert result["images"][0]["id"] == "img-0.jpeg"
    assert result["images"][0]["url"].endswith("/img-0.jpeg")
    assert uploaded == [{"user_id": "user-1", "image_id": "img-0.jpeg", "data": b"PNGDATA"}]


@pytest.mark.asyncio
async def test_document_parse_skips_images_when_disabled(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_tool

    _patch_common(monkeypatch)
    client = _FakeToolHttpClient([b"PDFBYTES"], _fake_ocr_payload())
    _patch_http(monkeypatch, client)
    uploaded = _patch_image_upload(monkeypatch)

    result = json.loads(
        await document_parse_tool.document_parse.coroutine(
            url="https://files.example.com/doc.pdf",
            include_images=False,
            runtime=_Runtime("user-1"),
        )
    )

    assert result["success"] is True
    assert "include_image_base64" not in client.post_calls[0]["json"]
    assert result["image_count"] == 0
    assert result["markdown"] == "# Hello\n\n![image](img-0.jpeg)\n\nSecond page"
    assert uploaded == []


@pytest.mark.asyncio
async def test_document_parse_rejects_unsupported_extension(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_tool

    _patch_common(monkeypatch)
    client = _FakeToolHttpClient([b"whatever"], _fake_ocr_payload())
    _patch_http(monkeypatch, client)

    result = json.loads(
        await document_parse_tool.document_parse.coroutine(
            url="https://files.example.com/sheet.xlsx",
            runtime=_Runtime("user-1"),
        )
    )

    assert "error" in result
    assert "xlsx" in result["error"]
    assert client.post_calls == []


@pytest.mark.asyncio
async def test_document_parse_returns_error_when_download_exceeds_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_tool

    _patch_common(monkeypatch)
    monkeypatch.setattr(document_parse_tool.settings, "DOCUMENT_PARSE_MAX_DOWNLOAD_BYTES", 8)
    client = _FakeToolHttpClient([b"0123456789abcdef"], _fake_ocr_payload())

    def _stream_with_headers(method: str, request_url: str):
        client.download_urls.append(request_url)
        return _FakeStreamResponse([b"0123456789abcdef"], {"content-length": "16"})

    client.stream = _stream_with_headers
    _patch_http(monkeypatch, client)

    result = json.loads(
        await document_parse_tool.document_parse.coroutine(
            url="https://files.example.com/doc.pdf",
            runtime=_Runtime("user-1"),
        )
    )

    assert "exceeds" in result["error"]
    assert client.post_calls == []


@pytest.mark.asyncio
async def test_document_parse_falls_back_to_markitdown_when_provider_fails(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_providers as providers
    from src.infra.tool import document_parse_tool

    _patch_common(monkeypatch)
    client = _FakeToolHttpClient([b"PDFBYTES"], RuntimeError("boom"))
    _patch_http(monkeypatch, client)

    async def _failing_parse(*args, **kwargs):
        raise providers.DocumentParseError("mistral HTTP 500: boom")

    monkeypatch.setitem(providers.PROVIDER_FUNCS, "mistral", _failing_parse)

    result = json.loads(
        await document_parse_tool.document_parse.coroutine(
            url="https://files.example.com/doc.pdf",
            runtime=_Runtime("user-1"),
        )
    )

    # mistral 失败后沿链降级到本地 markitdown 兜底
    assert result["success"] is True
    assert result["provider"] == "markitdown"


@pytest.mark.asyncio
async def test_document_parse_truncates_long_markdown(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_tool

    _patch_common(monkeypatch)
    monkeypatch.setattr(document_parse_tool.settings, "DOCUMENT_PARSE_MAX_OUTPUT_CHARS", 10)
    payload = {
        "model": "mistral-ocr-latest",
        "pages": [{"index": 1, "markdown": "a" * 50}],
        "usage_info": {"pages_processed": 1},
    }
    client = _FakeToolHttpClient([b"PDFBYTES"], payload)
    _patch_http(monkeypatch, client)

    result = json.loads(
        await document_parse_tool.document_parse.coroutine(
            url="https://files.example.com/doc.pdf",
            runtime=_Runtime("user-1"),
        )
    )

    assert result["success"] is True
    assert result["truncated"] is True
    assert len(result["markdown"]) <= 30


@pytest.mark.asyncio
async def test_upload_image_to_storage_uploads_under_user_folder(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from src.infra.tool import document_parse_tool

    async def fake_run_long_blocking_io(func, *args, **kwargs):
        return func(*args, **kwargs)

    monkeypatch.setattr(document_parse_tool, "run_long_blocking_io", fake_run_long_blocking_io)

    captured: dict[str, object] = {}

    class _FakeStorage:
        async def upload_file(self, file_obj, **kwargs):
            captured.update(kwargs)
            file_obj.seek(0)
            captured["content"] = file_obj.read()
            return SimpleNamespace(
                key=f"parsed-documents/u1/{kwargs['filename']}",
                url=f"https://storage.example.com/parsed-documents/u1/{kwargs['filename']}",
                size=len(captured["content"]),
                content_type=kwargs.get("content_type"),
            )

    async def fake_get_or_init_storage():
        return _FakeStorage()

    monkeypatch.setattr(
        "src.infra.storage.s3.service.get_or_init_storage", fake_get_or_init_storage
    )

    result = await document_parse_tool._upload_image_to_storage(
        user_id="u1", image_id="img-0.jpeg", data=b"IMG"
    )

    assert result is not None
    assert result["url"] == "https://storage.example.com/parsed-documents/u1/img-0.jpeg"
    assert result["size"] == 3
    assert result["content_type"] == "image/jpeg"
    assert captured["folder"] == "parsed-documents/u1"
    assert captured["filename"] == "img-0.jpeg"
    assert captured["skip_size_limit"] is True
    assert captured["content"] == b"IMG"


def test_get_document_parse_tool_returns_expected_tool() -> None:
    from src.infra.tool.document_parse_tool import get_document_parse_tool

    tool = get_document_parse_tool()

    assert tool.name == "document_parse"
